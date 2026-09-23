from datetime import datetime
from io import BytesIO

from django.contrib import admin
from django.contrib import messages
from django.http import HttpResponse, HttpResponseRedirect
from django.utils import timezone
from django.utils.safestring import mark_safe
from django.utils.translation import gettext_lazy as _
from django.urls import reverse
from django.db.models import Count
from django.db.models import IntegerField
from django.db.models import OuterRef
from django.db.models import Subquery
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from common.admin import FileInline
from common.utils.helpers import SAFE_SUBJECT_ICON
from common.utils.helpers import set_modified_by
from crm.models import ProductPriceTier
from crm.models.pricingtier import CategoryPriceRule
from crm.site.tieredpricinginline import ProductPriceTierInline
from common.utils.resize_image import resize_image
from crm.utils.admfilters import ByDepartmentFilter
from crm.utils.admfilters import ScrollRelatedOnlyFieldListFilter
from crm.models.pricingtier import DepartmentPriceRule

CATEGORY = _('Category')


class ProductAdmin(admin.ModelAdmin):
    actions = ('export_price_lists',)
    fieldsets = [
        (None, {
            'fields': (
                'main_image_preview', 'main_image',
                'name', 'type',
                'product_category',
                'description',
                # ('price', 'currency') not currently in use
            )
        }),
        (_('Additional information'), {
            'classes': ('collapse',),
            'fields': [
                'modified_by',
                ('creation_date', 'update_date'),
            ]
        }),
    ]
    inlines = [ProductPriceTierInline, FileInline]
    list_display = (
        'name_icon',
        'price',
        'currency',
        'get_type',
        'get_category'
    )
    list_filter = ('type', ('product_category',
                   ScrollRelatedOnlyFieldListFilter))
    radio_fields = {"type": admin.HORIZONTAL}
    readonly_fields = ('modified_by', 'creation_date', 'update_date',
                       'main_image_preview'
                       )
    save_on_top = True
    search_fields = ['name', 'description']

    # -- ModelAdmin methods -- #

    @admin.action(description=_('Export price lists to Excel'))
    def export_price_lists(self, request, queryset):
        price_tiers = (
            ProductPriceTier.objects
            .filter(product__in=queryset)
            .select_related(
                'product',
                'product__product_category',
                'product__department',
            )
            .order_by(
                'product__product_category__name',
                'product__name',
                '-price',
            )
        )

        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)

        sheets = {}
        used_titles = set()

        for price_tier in price_tiers:
            product = price_tier.product
            category = product.product_category
            category_name = category.name if category else _(
                'Without category')

            if category_name not in sheets:
                title = self._get_excel_sheet_title(
                    category_name,
                    used_titles,
                )
                sheets[category_name] = workbook.create_sheet(title)
                used_titles.add(title)

                self._format_price_sheet(sheets[category_name])

            # department = getattr(product, 'department', None)
            currency = getattr(
                product.department.department.default_currency, 'name', '')

            sheet = sheets[category_name]
            sheet.append([
                product.pk,
                product.name,
                price_tier.tier_name.name,
                price_tier.price,
                currency,
                price_tier.min_quantity,
                self._excel_datetime(price_tier.valid_from),
                self._excel_datetime(price_tier.valid_to),
                price_tier.description or '',
            ])

        if not sheets:
            sheet = workbook.create_sheet(_('Price lists')[:31])
            self._format_price_sheet(sheet)
            sheet.append([
                '',
                '',
                '',
                '',
                '',
                '',
                '',
                '',
                _('No price tiers found'),
            ])

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        filename = (
            f'price_lists_'
            f'{timezone.localdate().isoformat()}.xlsx'
        )

        response = HttpResponse(
            output.getvalue(),
            content_type=(
                'application/vnd.openxmlformats-officedocument'
                '.spreadsheetml.sheet'
            ),
        )
        response['Content-Disposition'] = (
            f'attachment; filename="{filename}"'
        )
        return response

    @staticmethod
    def _format_price_sheet(sheet):
        headers = [
            'product_id',
            'product_name',
            'tier_name',
            'price',
            'currency',
            'min_quantity',
            'valid_from',
            'valid_to',
            'description',
        ]

        sheet.append(headers)
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = 'A1:J1'
        sheet.sheet_view.showGridLines = False

        header_fill = PatternFill(
            fill_type='solid',
            fgColor='1F4E78',
        )

        for cell in sheet[1]:
            cell.font = Font(
                bold=True,
                color='FFFFFF',
            )
            cell.fill = header_fill
            cell.alignment = Alignment(
                horizontal='center',
                vertical='center',
            )

        widths = {
            'A': 12,
            'B': 32,
            'C': 24,
            'D': 14,
            'E': 12,
            'F': 14,
            'G': 16,
            'H': 16,
            'I': 45,
        }

        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

        sheet.row_dimensions[1].height = 30
        sheet.freeze_panes = 'A2'
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.fitToWidth = 1
        sheet.print_title_rows = '1:1'

    @staticmethod
    def _get_excel_sheet_title(category_name, used_titles):
        invalid_chars = '[]:*?/\\'
        title = ''.join(
            '_' if char in invalid_chars else char
            for char in str(category_name)
        ).strip()

        title = title or 'Without category'
        title = title[:31]

        base_title = title
        counter = 2

        while title in used_titles:
            suffix = f' ({counter})'
            title = f'{base_title[:31 - len(suffix)]}{suffix}'
            counter += 1

        return title

    @staticmethod
    def _excel_datetime(value):
        if value is None:
            return None

        if isinstance(value, datetime) and timezone.is_aware(value):
            return timezone.make_naive(value)

        return value

    def add_view(self, request, form_url='', extra_context=None):
        extra_context = extra_context or {}
        extra_context["show_add_price_tier"] = True
        return super().add_view(request, form_url, extra_context=extra_context)

    def change_view(self, request, object_id, form_url='', extra_context=None):
        extra_context = extra_context or {}
        counters = self._get_counters(object_id)
        if counters["price_rule"] == 0:
            extra_context["show_add_price_tier"] = True
        else:
            extra_context["show_add_price_tier"] = counters["price_tier"] < counters["price_rule"]
        return super().change_view(request, object_id, form_url, extra_context=extra_context)

    def changelist_view(self, request, extra_context=None):
        self.changelist_url = reverse("site:crm_product_changelist")
        self.query_dict = request.GET.copy()
        return super().changelist_view(
            request, extra_context=extra_context,
        )

    def get_list_filter(self, request):
        list_filter = list(self.list_filter)
        if request.user.is_superuser or request.user.is_chief:
            list_filter.insert(0, ByDepartmentFilter)
        return list_filter

    def get_queryset(self, request):
        qs = super().get_queryset(request)
        if request.user.department_id:
            qs = qs.filter(department_id=request.user.department_id)
        elif request.user.is_superoperator:
            qs = qs.filter(
                department__in=request.user.groups.filter(
                    department__isnull=False
                )
            )
        return qs

    def response_post_save_add(self, request, obj):
        if '_add-product-price-tier' not in request.POST:
            return super().response_post_save_add(request, obj)
        return self._create_missing_product_price_tiers(request, obj)

    def response_post_save_change(self, request, obj):
        if '_add-product-price-tier' not in request.POST:
            return super().response_post_save_change(request, obj)
        return self._create_missing_product_price_tiers(request, obj)

    def save_model(self, request, obj, form, change):
        # Resize new main image if present
        if 'main_image' in form.changed_data and obj.main_image:
            resize_main_image(obj)
        if change and 'main_image' in form.changed_data:
            # Delete old main image file if a new one is being uploaded
            try:
                old_instance = self.model.objects.get(pk=obj.pk)
                if old_instance.main_image and old_instance.main_image != obj.main_image:
                    old_instance.main_image.delete(save=False)
            except self.model.DoesNotExist:
                pass
        obj.modified_by = request.user
        if not obj.department and request.user.department_id:
            obj.department_id = request.user.department_id
        super().save_model(request, obj, form, change)

    def save_related(self, request, form, formsets, change):
        # Save who modified the product price tier instances.
        set_modified_by(request, formsets, ProductPriceTier)
        super().save_related(request, form, formsets, change)

    # -- ModelAdmin callables -- #

    @admin.display(
        description=CATEGORY,
        ordering='product_category__name'
    )
    def get_category(self, obj):
        category = obj.product_category
        if category:
            self.query_dict['product_category__id__exact'] = category.id
            url = f"{self.changelist_url}?{self.query_dict.urlencode()}"
            return mark_safe(
                f'<a href="{url}" title="{CATEGORY}">{category.name}</a>'
            )

    @admin.display(
        description=_("Type"),
        ordering='type'
    )
    def get_type(self, obj):
        self.query_dict['type__exact'] = obj.type
        url = f"{self.changelist_url}?{self.query_dict.urlencode()}"
        return mark_safe(
            f'<a href="{url}" title="{_('Type')}">{obj.TYPE_CHOICES[obj.type]}</a>'
        )

    @admin.display(description='')
    def main_image_preview(self, obj):
        if obj.main_image:
            return mark_safe(
                f'<img src="{obj.main_image.url}" style="width:200px;height:200px;">'
            )
        return mark_safe(
            '<i class="material-icons" style="font-size: 200px;vertical-align: middle;'
            'border-radius:50%;color: var(--body-quiet-color)">inventory_2</i>'
        )

    @admin.display(description=SAFE_SUBJECT_ICON, ordering='name')
    def name_icon(self, obj):
        if obj.main_image:
            return mark_safe(
                f'<span style="white-space: nowrap;">'
                f'<img src="{obj.main_image.url}" style="vertical-align: middle;'
                'width:20px;height:20px;">'
                f'&nbsp;{obj.name}</span>'
            )
        return obj.name

# -- Custom methods -- #

    def _create_missing_product_price_tiers(self, request, obj) -> HttpResponseRedirect:
        """Create missing product price tiers from category or department price rules."""
        category_rules = CategoryPriceRule.objects.filter(
            category=obj.product_category
        ) if obj.product_category_id else CategoryPriceRule.objects.none()

        # Category price rules have priority over department rules.
        price_rules = category_rules if category_rules.exists() else (
            DepartmentPriceRule.objects.filter(
                department=obj.department
            )
        )

        existing_price_tiers = ProductPriceTier.objects.filter(product=obj)
        existing_names = set(
            existing_price_tiers.values_list('tier_name', flat=True))

        base_rule = price_rules.filter(base_tier=True).first()
        base_price_tier = existing_price_tiers.filter(
            tier_name=base_rule.tier_name
        ).first() if base_rule else None

        if not base_price_tier:
            if base_rule:
                ProductPriceTier.objects.create(
                    tier_name=base_rule.tier_name,
                    price=0.00,
                    product=obj,
                    department=obj.department,
                    modified_by=request.user,
                )
                messages.warning(
                    request,
                    _("Specify the price value in the created price tier."),
                )
            else:
                messages.error(
                    request,
                    _("First, add rules for creating price tiers."),
                )
        elif not base_price_tier.price:
            messages.warning(
                request,
                _("The base price tier must have a price value before adding other tiers."),
            )
        else:
            for price_rule in price_rules:
                if price_rule.tier_name_id in existing_names:
                    continue

                if price_rule.price_type == 'D':
                    price = base_price_tier.price * (
                        1 - price_rule.percentage / 100
                    )
                elif price_rule.price_type == 'M':
                    price = base_price_tier.price * (
                        1 + price_rule.percentage / 100
                    )
                else:
                    price = 0.00
                    messages.warning(
                        request,
                        _("Specify the price value in the created price tier."),
                    )

                ProductPriceTier.objects.create(
                    tier_name=price_rule.tier_name,
                    price=price,
                    product=obj,
                    department=obj.department,
                    modified_by=request.user,
                )
        if 'admin' in request.path:
            url = reverse("admin:crm_product_change", args=(obj.id,))
        else:
            url = reverse("site:crm_product_change", args=(obj.id,))
        return HttpResponseRedirect(
            f"{url}?{request.META.get('QUERY_STRING', '')}"
        )

    def _get_counters(self, object_id):
        dept_rules_sq = DepartmentPriceRule.objects.filter(
            department=OuterRef('department')
        ).values('department').annotate(
            n=Count('pk')
        ).values('n')

        category_rules_sq = CategoryPriceRule.objects.filter(
            category=OuterRef('product_category')
        ).values('category').annotate(
            n=Count('pk')
        ).values('n')

        product_price_tiers_sq = ProductPriceTier.objects.filter(
            product=OuterRef('pk')
        ).values('product').annotate(
            n=Count('pk')
        ).values('n')

        product = self.model.objects.annotate(
            department_price_rule=Subquery(
                dept_rules_sq,
                output_field=IntegerField(),
            ),
            category_price_rule=Subquery(
                category_rules_sq,
                output_field=IntegerField(),
            ),
            product_price_tier=Subquery(
                product_price_tiers_sq,
                output_field=IntegerField(),
            ),
        ).only('pk', 'department', 'product_category').get(pk=object_id)

        category_count = product.category_price_rule or 0

        return {
            "price_rule": (
                category_count
                or product.department_price_rule
                or 0
            ),
            "price_tier": product.product_price_tier or 0,
        }

# -- Custom functions -- #


def resize_main_image(obj) -> None:
    """
    Resize uploaded main image to a maximum of 200x200 pixels.
    """
    if obj.main_image:
        resized_image = resize_image(obj.main_image)

        # Create a new File object
        obj.main_image.file = resized_image
        obj.main_image.name = f"{obj.name}.png"
